"""Check the delivered workbook and actual Excel recalculation evidence."""
import pytest
import pandas as pd
from openpyxl import load_workbook
from src.config import ROOT

@pytest.mark.integration
def test_excel_formula_results():
    p=ROOT/'outputs/aero-release/Motorsport_Aero_Engineering.xlsx'
    if not p.exists():pytest.skip('Workbook not generated')
    book=load_workbook(p,data_only=True,read_only=True)
    assert len(book.sheetnames)==31
    for row in book['Python Reconciliation'].iter_rows(min_row=6,max_row=10,values_only=True):
        assert row[2] is not None and abs(row[3])<1e-6
    current=pd.read_csv(ROOT/'data/processed/recommended_trace_2.csv').dt_s.sum()
    assert book['Python Reconciliation']['B10'].value==pytest.approx(current,abs=1e-6)
    book.close()

@pytest.mark.integration
def test_excel_formula_and_input_protection():
    p=ROOT/'outputs/aero-release/Motorsport_Aero_Engineering.xlsx'
    if not p.exists():pytest.skip('Workbook not generated')
    book=load_workbook(p,data_only=False)
    assert book['Aero Calculations']['C16'].value.startswith('=')
    assert book['Aero Calculations'].protection.sheet
    assert not book['Assumptions']['B6'].protection.locked
    assert len(book['Charts']._charts)>=1
    assert len(book['Assumptions'].data_validations.dataValidation)>=1
    book.close()
